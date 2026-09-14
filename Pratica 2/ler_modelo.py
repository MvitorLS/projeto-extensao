from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document(r'c:\Users\Vitor\Documents\projeto extensão\modelo-relatorio-final.docx')

print("=== ESTILOS USADOS ===")
estilos = set()
for p in doc.paragraphs:
    estilos.add(p.style.name)
for s in sorted(estilos):
    print(f"  {s}")

print("\n=== SEÇÕES / MARGENS ===")
for i, sec in enumerate(doc.sections):
    print(f"Seção {i}: top={sec.top_margin.cm:.1f}cm bottom={sec.bottom_margin.cm:.1f}cm left={sec.left_margin.cm:.1f}cm right={sec.right_margin.cm:.1f}cm")

print("\n=== PARÁGRAFOS (primeiros 80) ===")
for i, p in enumerate(doc.paragraphs[:80]):
    runs_info = []
    for r in p.runs:
        info = []
        if r.bold: info.append("BOLD")
        if r.italic: info.append("ITALIC")
        if r.font.size: info.append(f"{r.font.size.pt:.0f}pt")
        if r.font.name: info.append(r.font.name)
        runs_info.append(f"[{','.join(info)}]" if info else "[]")

    align = str(p.alignment).replace("WD_PARAGRAPH_ALIGNMENT.", "") if p.alignment else "None"
    indent = f"first={p.paragraph_format.first_line_indent}" if p.paragraph_format.first_line_indent else ""
    space_b = f"sb={p.paragraph_format.space_before.pt:.0f}" if p.paragraph_format.space_before else ""
    space_a = f"sa={p.paragraph_format.space_after.pt:.0f}" if p.paragraph_format.space_after else ""

    text_preview = p.text[:80].replace('\n', '↵')
    print(f"[{i:3d}] style='{p.style.name}' align={align} {indent} {space_b} {space_a}")
    print(f"       runs={runs_info}")
    print(f"       text='{text_preview}'")

print("\n=== TABELAS ===")
for i, t in enumerate(doc.tables):
    print(f"Tabela {i}: {len(t.rows)} linhas x {len(t.columns)} colunas")
    for r in t.rows[:2]:
        row_text = [c.text[:30] for c in r.cells]
        print(f"  {row_text}")
