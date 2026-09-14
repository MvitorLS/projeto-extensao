from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document(r'c:\Users\Vitor\Documents\projeto extensão\modelo-relatorio-final.docx')

print("=== PARÁGRAFOS COMPLETOS ===")
for i, p in enumerate(doc.paragraphs):
    runs_info = []
    for r in p.runs:
        info = []
        if r.bold: info.append("BOLD")
        if r.italic: info.append("ITALIC")
        try:
            if r.font.size: info.append(f"{r.font.size.pt:.0f}pt")
        except: pass
        if r.font.name: info.append(r.font.name)
        runs_info.append(f"[{','.join(info)}]" if info else "[]")

    align = str(p.alignment).replace("WD_PARAGRAPH_ALIGNMENT.", "") if p.alignment else "None"

    try:
        fi = p.paragraph_format.first_line_indent
        indent = f"fi={fi.cm:.2f}cm" if fi else ""
    except:
        indent = "fi=?"

    try:
        sb = p.paragraph_format.space_before
        sb_str = f"sb={sb.pt:.0f}pt" if sb else ""
    except:
        sb_str = ""

    try:
        sa = p.paragraph_format.space_after
        sa_str = f"sa={sa.pt:.0f}pt" if sa else ""
    except:
        sa_str = ""

    try:
        li = p.paragraph_format.left_indent
        li_str = f"li={li.cm:.2f}cm" if li else ""
    except:
        li_str = ""

    text = p.text[:100].replace('\n', '↵')
    print(f"[{i:3d}] align={align:10s} {indent:12s} {sb_str:8s} {sa_str:8s} {li_str:10s} | {text[:70]}")

print("\n=== LINHA SPACING de parágrafos com texto ===")
for i, p in enumerate(doc.paragraphs):
    if not p.text.strip():
        continue
    try:
        ls = p.paragraph_format.line_spacing
        ls_rule = p.paragraph_format.line_spacing_rule
        print(f"[{i:3d}] ls={ls} rule={ls_rule} | {p.text[:60]}")
    except:
        pass
